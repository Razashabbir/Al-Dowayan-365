"""Read the account -> statement-line mapping from an "FS ... Mapping" workbook (sheet "TB25 Dec" or "TB Dec").

Columns used (as in the YE2025 workbooks): A account, B ledger account description, C FSLI,
M FSLI (final), N note line.  Rows whose account is not a plain number (PwC1, 4216001.1 ...) are audit
adjustments that do not exist in D365 and are skipped.

    python fs_import.py "D:\\games\\Al Dowayan\\FS Tazaiud Mapping YE2025 (DEC) 1.xlsx" Tazayud
writes the mapping straight into rpt.fs_map (set name = second argument).
"""
import re
import sys

import fs_model

SHEETS = ("TB25 Dec", "TB Dec")


def read_workbook(path: str) -> list[dict]:
    import openpyxl   # pip install openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next((wb[s] for s in SHEETS if s in wb.sheetnames), None)
    if sheet is None:
        raise ValueError(f"No sheet named {' or '.join(SHEETS)} in {path}")
    out, seen = [], set()
    for r in sheet.iter_rows(min_row=5, values_only=True):
        r = list(r) + [None] * 16
        acc = r[0]
        if isinstance(acc, float) and acc.is_integer():
            acc = int(acc)
        acc = str(acc).strip() if acc is not None else ""
        if not re.fullmatch(r"\d{5,10}", acc) or acc in seen:
            continue
        seen.add(acc)
        name = str(r[1] or "").strip()
        fsli = str(r[12] or r[2] or "").strip()
        note = str(r[13] or "").strip()
        out.append({"main_account": acc, "account_name": name[:200], "fsli": fsli[:160], "note_line": note[:160],
                    "line_code": fs_model.line_from_workbook(acc, name, fsli, note)})
    return out


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    import os
    sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "etl"))
    from dotenv import load_dotenv
    load_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "etl", ".env"))
    import db
    from sqlalchemy import text
    rows = read_workbook(sys.argv[1])
    with db.engine().begin() as cn:
        cn.execute(text("DELETE FROM rpt.fs_map WHERE map_set = :s AND tenant_key IS NULL"), {"s": sys.argv[2]})
        for e in rows:
            cn.execute(text("""INSERT INTO rpt.fs_map (map_set, tenant_key, main_account, account_name, fsli, note_line, line_code)
                               VALUES (:s, NULL, :main_account, :account_name, :fsli, :note_line, :line_code)"""),
                       {"s": sys.argv[2], **e})
    print(f"{len(rows)} accounts imported into mapping set '{sys.argv[2]}'.")
